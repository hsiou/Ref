# -*- coding: utf-8 -*-
"""按论文.xlsx 模板生成两个结果副本（全部结果 / 2024-01-01 之后）。

输入：
  outputs/pubmed_verified.json   Step4 PubMed 核验结果
  outputs/cnki_results.json      Step3 CNKI 检索结果（标题/作者/来源/日期）
输出：
  outputs/论文_全部结果.xlsx
  outputs/论文_2024年之后.xlsx
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
TEMPLATE = Path(r"E:\Download\论文.xlsx")
HEADER = (
    "序号",
    "成果类型",
    "成果名称",
    "完成人",
    "刊物名称及影响因子、出版社或授权单位名称",
    "年、卷、期、页/或专利号",
    "他引次数/或出版数量",
)


def load_pubmed() -> list[dict]:
    recs = json.loads((OUT / "pubmed_verified.json").read_text(encoding="utf-8"))
    rows = []
    for r in recs:
        authors = "，".join(a.get("display_name") or "" for a in r.get("authors", []))
        vol, iss, pages = r.get("volume") or "", r.get("issue") or "", r.get("pages") or ""
        parts = [r.get("year") or ""]
        if vol or iss:
            parts.append(f"{vol}({iss})" if vol and iss else (vol or iss))
        if pages:
            parts.append(pages)
        rows.append(
            {
                "type": "期刊论文",
                "title": r.get("title") or "",
                "authors": authors,
                "source": r.get("journal") or "",
                "bib": "，".join(p for p in parts if p),
                "year": int(r["year"]) if str(r.get("year", "")).isdigit() else None,
                "source_db": "PubMed",
            }
        )
    return rows


def load_cnki() -> list[dict]:
    path = OUT / "cnki_results.json"
    if not path.exists():
        return []
    recs = json.loads(path.read_text(encoding="utf-8"))
    rows = []
    for r in recs:
        date = str(r.get("date") or "")
        m = re.match(r"(\d{4})", date)
        rows.append(
            {
                "type": "期刊论文",
                "title": r.get("title") or "",
                "authors": "，".join(r.get("authors") or []),
                "source": r.get("journal") or "",
                "bib": date,
                "year": int(m.group(1)) if m else None,
                "source_db": "CNKI",
            }
        )
    return rows


def write_book(rows: list[dict], dest: Path) -> None:
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    # 保留模板表头，清空可能存在的旧数据行
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for i, r in enumerate(rows, start=1):
        ws.append(
            [
                i,
                r["type"],
                r["title"],
                r["authors"],
                r["source"],
                r["bib"],
                "",  # 他引次数模板列留空
            ]
        )
    wb.save(dest)
    print(f"wrote {dest.name}: {len(rows)} rows")


def main() -> None:
    rows = load_pubmed() + load_cnki()
    # 全部结果：PubMed 在前、CNKI 在后；2024 副本按发表年份过滤
    write_book(rows, OUT / "论文_全部结果.xlsx")
    after = [r for r in rows if r["year"] is not None and r["year"] >= 2024]
    write_book(after, OUT / "论文_2024年之后.xlsx")


if __name__ == "__main__":
    main()
